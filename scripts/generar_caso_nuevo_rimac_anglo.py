from pathlib import Path
from zipfile import ZipFile
import shutil
from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
PLANTILLA = BASE_DIR / "plantillas_maestras" / "plantilla_maestra_clean_9paginas.docx"
OUT_DOCX = BASE_DIR / "generados" / "RESOLUCION_0842-2026_CC1_RIMAC_ANGLO.docx"
OUT_PDF = BASE_DIR / "generados" / "RESOLUCION_0842-2026_CC1_RIMAC_ANGLO.pdf"
CAPTURA_DIR = BASE_DIR / "generados" / "capturas_nueva_resolucion"
CAPTURA_DIR.mkdir(parents=True, exist_ok=True)
ARTIFACT_DIR = Path(r"C:\Users\Admin\.gemini\antigravity\brain\18369d91-b2f7-445f-b6d3-baa02401afcd")

def set_paragraph_text_preserving_footnotes(p, new_text_before_fn=None, new_text_after_fn=None, full_text=None):
    fn_indices = [idx for idx, r in enumerate(p.runs) if 'footnoteReference' in r._r.xml]
    if not fn_indices:
        if full_text is not None:
            if p.runs:
                p.runs[0].text = full_text
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = full_text
        return
        
    fn_idx = fn_indices[0]
    if new_text_before_fn is not None:
        for i in range(fn_idx):
            if i == 0:
                p.runs[i].text = new_text_before_fn
            else:
                p.runs[i].text = ""
                
    if new_text_after_fn is not None:
        after_runs = list(range(fn_idx + 1, len(p.runs)))
        if after_runs:
            p.runs[after_runs[0]].text = new_text_after_fn
            for i in after_runs[1:]:
                p.runs[i].text = ""
        else:
            p.add_run(new_text_after_fn)

def generar_nueva_resolucion():
    shutil.copyfile(PLANTILLA, OUT_DOCX)
    doc = Document(OUT_DOCX)
    
    # 1. P0 a P3: Metadatos de control
    set_paragraph_text_preserving_footnotes(doc.paragraphs[0], full_text="Elaborado por:\tDavid Chávez")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[1], full_text="Supervisado por:\tLoussiana Salazar")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[2], full_text="Equipo:\tSeguros y Salud")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[3], full_text="Vencimiento:\t28/11/2026")
    
    # P5: Título de la Resolución
    set_paragraph_text_preserving_footnotes(doc.paragraphs[5], full_text="RESOLUCIÓN FINAL N° 0842-2026/CC1")
    
    # P7: Denunciante
    set_paragraph_text_preserving_footnotes(doc.paragraphs[7], full_text="DENUNCIANTE\t:\tROSA ELVIRA MENDOZA CARRIÓN (SEÑORA MENDOZA)")
    
    # P8: Denunciados
    set_paragraph_text_preserving_footnotes(doc.paragraphs[8], full_text="DENUNCIADOS\t:\tRÍMAC SEGUROS Y REASEGUROS S.A. (RÍMAC SEGUROS)\n\t\t\tCLÍNICA ANGLO AMERICANA S.A. (CLÍNICA)")
    
    # P9 a P11: Materia y Actividad
    set_paragraph_text_preserving_footnotes(doc.paragraphs[9], full_text="MATERIA\t:\tIMPROCEDENCIA DE LA DENUNCIA")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[10], full_text="\t\tDECLINACIÓN DE COMPETENCIA")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[11], full_text="ACTIVIDAD\t:\tACTIVIDADES RELACIONADAS CON LA SALUD HUMANA")
    
    # P13: Fecha
    set_paragraph_text_preserving_footnotes(doc.paragraphs[13], full_text="Lima, 28 de setiembre de 2026")
    
    # P15: Título ANTECEDENTES
    set_paragraph_text_preserving_footnotes(doc.paragraphs[15], full_text="ANTECEDENTES")
    
    # P17: Introducción de antecedentes (CON FOOTNOTE 1)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[17],
        new_text_before_fn="Mediante escrito del 18 de mayo de 2026, subsanado mediante escrito del 12 de junio de 2026, la señora Mendoza denunció a Rímac Seguros y a la Clínica por presuntas infracciones a la Ley N° 29571, Código de Protección y Defensa del Consumidor",
        new_text_after_fn=" (en adelante, el Código), señalando lo siguiente:"
    )
    
    # Hechos fácticos numerados (P19 a P31)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[19],
        full_text="El 12 de febrero de 2026, la denunciante contaba con la Póliza de Asistencia Médica y Salud Integral N° 09283741 emitida por Rímac Seguros, encontrándose con las primas comerciales debidamente canceladas y al día en sus pagos."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[21],
        full_text="El 24 de abril de 2026, presentó un cuadro agudo de dolor abdominal severo, siendo ingresada en estado de emergencia a la Clínica, donde se le diagnosticó apendicitis aguda complicada con peritonitis y requirió intervención quirúrgica de urgencia mediante técnica laparoscópica."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[23],
        full_text="Rímac Seguros se negó injustificadamente a otorgar la cobertura médica de emergencia requerida para la cirugía y hospitalización, alegando erróneamente la preexistencia de patologías gastrointestinales no acreditadas."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[25],
        full_text="Por su parte, la Clínica emitió una preliquidación y liquidación exigiendo el cobro no presupuestado de S/ 9 450,00 por concepto de instrumental quirúrgico laparoscópico e insumos descartables no informados previamente a la paciente."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[27],
        full_text="Asimismo, al solicitar el alta médica, la Clínica se negó a hacer entrega de la copia completa y fedateada de la historia clínica, informes patológicos y epicrisis, condicionando su salida al pago inmediato o entrega de un pagaré en garantía."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[29],
        full_text="Pese a los reiterados reclamos formales cursados por la señora Mendoza, Rímac Seguros mantuvo la denegatoria de cobertura y la Clínica persistió en la retención arbitraria de su acervo clínico documentario."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[31],
        full_text="Por tales motivos, interpuso denuncia administrativa ante el Indecopi requiriendo la imposición de sanciones y medidas correctivas reparadoras a los proveedores denunciados."
    )
    
    # P33: Medidas correctivas solicitadas
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[33],
        full_text="La señora Mendoza solicitó, en calidad de medida correctiva, que: (i) Rímac Seguros cumpla con otorgar la cobertura integral de gastos médicos y quirúrgicos derivados de la emergencia; (ii) la Clínica cumpla con devolver el cobro indebido de S/ 9 450,00 y haga entrega inmediata de la copia fedateada de su historia clínica; y, (iii) se declare la nulidad de las garantías exigidas. Asimismo, solicitó la condena en costas y costos."
    )
    
    # P37: Subtítulo ANÁLISIS
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[37],
        full_text="Sobre la improcedencia de la denuncia interpuesta por la señora Mendoza"
    )
    
    # P45: Verificación de competencia
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[45],
        full_text="Al respecto, la Comisión de Protección al Consumidor N° 1 (en adelante, la Comisión) debe verificar si el Indecopi es el órgano competente para pronunciarse sobre los hechos cuestionados por la señora Mendoza en su denuncia contra Rímac Seguros y la Clínica."
    )
    
    # P59: Subsunción normativa de IAFAS + IPRESS (CON FOOTNOTE 9)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[59],
        new_text_before_fn="Asimismo, conforme al Anexo I-C del Reglamento de Infracciones y Sanciones de Susalud, aprobado por el Decreto Supremo N° 031-2014-SA",
        new_text_after_fn=", dicha entidad se encuentra facultada para sancionar el hecho de no brindar cobertura oportuna a los afiliados o beneficiarios respecto de las IAFAS; y, conforme al Anexo I-B del citado cuerpo normativo, sancionar a las IPRESS por realizar cobros no pactados, deficiencias asistenciales o negativa de entrega de documentación médica obligatoria."
    )
    
    # P77: Determinación de autoridad competente
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[77],
        full_text="A continuación, se determinará la autoridad competente para conocer la denuncia interpuesta por la señora Mendoza contra Rímac Seguros y la Clínica, de ser el caso, imponer las sanciones correspondientes."
    )
    
    # P85: Subsunción fáctica al caso concreto (calibrada exactamente para encaje de 9 páginas)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[85],
        full_text="En el presente caso, la señora Mendoza cuestionó que: (i) Rímac Seguros (IAFAS) no brindó oportunamente la cobertura de gastos médicos en situación de emergencia; (ii) la Clínica (IPRESS) exigió un cobro excesivo no pactado de S/ 9 450,00 por material laparoscópico; y, (iii) la Clínica se negó a entregar su historia clínica condicionando el alta. Dichas conductas inciden directamente en la cobertura del aseguramiento y en las prestaciones asistenciales de salud."
    )
    
    # P87: Conclusión de improcedencia
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[87],
        full_text="En consecuencia, la Comisión considera que corresponde declarar improcedente la denuncia interpuesta por la señora Mendoza contra Rímac Seguros y la Clínica, en la medida que las conductas cuestionadas resultan ser materia de competencia exclusiva de Susalud."
    )
    
    # P93: RESUELVE PRIMERO - SOLO 'PRIMERO: ' EN NEGRITA
    p93 = doc.paragraphs[93]
    p93.runs[0].text = "PRIMERO: "
    p93.runs[0].bold = True
    p93.runs[1].text = (
        "declarar improcedente la denuncia interpuesta por la señora Rosa Elvira Mendoza Carrión "
        "contra Rímac Seguros y Reaseguros S.A. y Clínica Anglo Americana S.A., por presunta infracción "
        "a la Ley N° 29571, Código de Protección y Defensa del Consumidor, en la medida que ha quedado "
        "acreditado que las conductas cuestionadas resultan ser materia de exclusiva competencia de la "
        "Superintendencia Nacional de Salud. En consecuencia, disponer la devolución a la parte denunciante de "
        "la tasa por derecho de trámite pagada"
    )
    p93.runs[1].bold = False
    for i in range(2, 16):
        p93.runs[i].text = ""
        p93.runs[i].bold = False
    p93.runs[17].text = ", previa solicitud a la Unidad de Finanzas y Contabilidad del Indecopi."
    p93.runs[17].bold = False
    for i in range(18, len(p93.runs)):
        p93.runs[i].text = ""
        p93.runs[i].bold = False
        
    # P95: RESUELVE SEGUNDO - SOLO 'SEGUNDO:' EN NEGRITA
    p95 = doc.paragraphs[95]
    p95.runs[0].text = "SEGUNDO:"
    p95.runs[0].bold = True
    p95.runs[1].text = " ordenar a la Secretaría Técnica de la Comisión de Protección al Consumidor N° 1 que remita "
    p95.runs[1].bold = False
    p95.runs[2].text = "el original"
    p95.runs[2].bold = False
    p95.runs[3].text = " de todo lo actuado en el presente procedimiento a la Superintendencia Nacional de Salud, a efectos de que adopte las medidas correspondientes en el ámbito de su competencia."
    p95.runs[3].bold = False
    
    # P97: RESUELVE TERCERO - SOLO 'TERCERO: ' EN NEGRITA
    p97 = doc.paragraphs[97]
    p97.runs[0].text = "TERCERO: "
    p97.runs[0].bold = True
    p97.runs[1].text = "informar "
    p97.runs[1].bold = False
    p97.runs[2].text = "a la "
    p97.runs[2].bold = False
    p97.runs[3].text = ""
    p97.runs[4].text = "señora "
    p97.runs[4].bold = False
    p97.runs[5].text = ""
    p97.runs[6].text = ""
    p97.runs[7].text = "Rosa Elvira Mendoza Carrión"
    p97.runs[7].bold = False
    
    # P98 y P99: Comisionados y Presidenta
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[98],
        full_text="Con la intervención de los señores Comisionados: Mónica Tatiana Siverio Puycan, María de Fátima Ponce Regalado, Ernesto Alonso Calderón Burneo y Aldrin Capcha Coronado."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[99],
        full_text="MÓNICA TATIANA SIVERIO PUYCAN\nPresidenta"
    )
    
    doc.save(OUT_DOCX)
    print("DOCX nuevo guardado con éxito.")
    
    # 2. Actualizar encabezado header1.xml con Expediente 0842-2026/CC1
    import tempfile
    temp_zip = Path(tempfile.mktemp(suffix=".docx"))
    with ZipFile(OUT_DOCX, "r") as zin:
        with ZipFile(temp_zip, "w") as zout:
            for item in zin.infolist():
                if item.filename == "word/header1.xml":
                    hdr = zin.read(item.filename).decode("utf-8")
                    hdr = hdr.replace("2685-2025/CC1", "0842-2026/CC1").replace("0512-2026/CC1", "0842-2026/CC1")
                    zout.writestr(item, hdr.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))
    shutil.move(temp_zip, OUT_DOCX)
    print("Header actualizado con Expediente 0842-2026/CC1.")

if __name__ == "__main__":
    generar_nueva_resolucion()

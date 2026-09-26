"""Guardia DLP (Data Loss Prevention / Fuga de Placeholders) para Resoluciones de Improcedencia.
Detecta cadenas parásitas, variables sin resolver o corchetes residuales.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import List, Tuple

from docx import Document

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

PATRONES_FUGA = [
    (r"\[NOMBRE\]", re.IGNORECASE),
    (r"\[EXPEDIENTE\]", re.IGNORECASE),
    (r"\[DENUNCIANTE\]", re.IGNORECASE),
    (r"\[DENUNCIADO\]", re.IGNORECASE),
    (r"\[FECHA\]", re.IGNORECASE),
    (r"\[XXX+\]", re.IGNORECASE),
    (r"\{[A-Za-z0-9_]+\}", re.IGNORECASE),
    (r"\bTODO\b", 0),  # En mayúsculas exactas para no colisionar con 'todo lo actuado' en español
    (r"\bFIXME\b", re.IGNORECASE),
    (r"\bCOMPLETAR\b", re.IGNORECASE),
    (r"\bINSERTAR\b", re.IGNORECASE),
]


def auditar_documento(ruta_docx: Path) -> List[Tuple[str, str, int]]:
    """Inspecciona párrafos y tablas buscando placeholders sin resolver."""
    doc = Document(str(ruta_docx))
    hallazgos = []

    for idx, p in enumerate(doc.paragraphs, 1):
        texto = p.text
        for pat, flags in PATRONES_FUGA:
            if re.search(pat, texto, flags):
                hallazgos.append((pat, texto[:100], idx))

    for t_idx, tabla in enumerate(doc.tables, 1):
        for f_idx, fila in enumerate(tabla.rows, 1):
            for c_idx, celda in enumerate(fila.cells, 1):
                texto = celda.text
                for pat, flags in PATRONES_FUGA:
                    if re.search(pat, texto, flags):
                        hallazgos.append((pat, f"Tabla {t_idx} Fila {f_idx} Celda {c_idx}: {texto[:60]}", 0))

    return hallazgos


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python guardia_improcedencia.py <archivo.docx>")
        sys.exit(1)

    archivo = Path(sys.argv[1])
    if not archivo.exists():
        print(f"Error: No existe el archivo {archivo}")
        sys.exit(1)

    fugas = auditar_documento(archivo)
    if fugas:
        print(f"❌ FALLO DE GUARDIA DLP: Se encontraron {len(fugas)} elementos sin resolver:")
        for pat, texto, parrafo in fugas:
            print(f"  - Patrón '{pat}' en párrafo/bloque {parrafo}: {texto}")
        sys.exit(1)
    else:
        print("✅ GUARDIA DLP APROBADA: Cero placeholders o variables sin resolver.")

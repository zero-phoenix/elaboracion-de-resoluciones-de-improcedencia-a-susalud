from pathlib import Path
from zipfile import ZipFile
import shutil
import re
from docx import Document

SRC_DOCX = Path(r"C:\Users\Admin\Desktop\impros susalud\MODELOS IMPROS SUSALUD CON FUNDAMENTO MEJORADO\2685-2025 RXX IMPRO SUSALUD okOK .docx")
BASE_DIR = Path(r"c:\Users\Admin\Documents\antigravity\zealous-kepler\elaboracion-de-resoluciones-de-improcedencia-a-susalud")
OUT_DOCX = BASE_DIR / "generados" / "RESOLUCION_0512-2026_CC1_MIXTA_COMPLEJA.docx"

def set_paragraph_text_preserving_footnotes(p, new_text_before_fn=None, new_text_after_fn=None, full_text=None):
    """
    Modifica el texto de un párrafo preservando intactos los runs que contienen <w:footnoteReference>.
    """
    fn_indices = [idx for idx, r in enumerate(p.runs) if 'footnoteReference' in r._r.xml]
    
    if not fn_indices:
        # No tiene footnote: podemos reemplazar todo el texto directamente
        if full_text is not None:
            # Mantener el primer run para preservar formato y vaciar los demás
            if p.runs:
                p.runs[0].text = full_text
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = full_text
        return
        
    fn_idx = fn_indices[0]
    
    # Hay footnote: preservar p.runs[fn_idx]
    if new_text_before_fn is not None:
        # Texto antes de la footnote
        for i in range(fn_idx):
            if i == 0:
                p.runs[i].text = new_text_before_fn
            else:
                p.runs[i].text = ""
                
    if new_text_after_fn is not None:
        # Texto después de la footnote
        after_runs = list(range(fn_idx + 1, len(p.runs)))
        if after_runs:
            p.runs[after_runs[0]].text = new_text_after_fn
            for i in after_runs[1:]:
                p.runs[i].text = ""
        else:
            p.add_run(new_text_after_fn)

def generate_mixta():
    shutil.copyfile(SRC_DOCX, OUT_DOCX)
    doc = Document(OUT_DOCX)
    
    # P0 a P3: Metadatos de control
    set_paragraph_text_preserving_footnotes(doc.paragraphs[0], full_text="Elaborado por:\tDavid Chávez")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[1], full_text="Supervisado por:\tLoussiana Salazar")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[2], full_text="Equipo:\tSeguros y Salud")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[3], full_text="Vencimiento:\t28/10/2026")
    
    # P5: Título de la Resolución
    set_paragraph_text_preserving_footnotes(doc.paragraphs[5], full_text="RESOLUCIÓN FINAL N° 0512-2026/CC1")
    
    # P7: Denunciante
    set_paragraph_text_preserving_footnotes(doc.paragraphs[7], full_text="DENUNCIANTE\t:\tMARCO ANTONIO REYES SALVATIERRA (SEÑOR REYES)")
    
    # P8: Denunciados
    set_paragraph_text_preserving_footnotes(doc.paragraphs[8], full_text="DENUNCIADOS\t:\tPACÍFICO COMPAÑÍA DE SEGUROS Y REASEGUROS S.A. (PACÍFICO SEGUROS)\n\t\t\tCLÍNICA SAN FELIPE S.A. (CLÍNICA)")
    
    # P9 a P11: Materia y Actividad
    set_paragraph_text_preserving_footnotes(doc.paragraphs[9], full_text="MATERIA\t:\tIMPROCEDENCIA DE LA DENUNCIA")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[10], full_text="\t\tDECLINACIÓN DE COMPETENCIA")
    set_paragraph_text_preserving_footnotes(doc.paragraphs[11], full_text="ACTIVIDAD\t:\tACTIVIDADES RELACIONADAS CON LA SALUD HUMANA")
    
    # P13: Fecha
    set_paragraph_text_preserving_footnotes(doc.paragraphs[13], full_text="Lima, 25 de setiembre de 2026")
    
    # P15: Título ANTECEDENTES
    set_paragraph_text_preserving_footnotes(doc.paragraphs[15], full_text="ANTECEDENTES")
    
    # P17: Introducción de antecedentes (TIENE FOOTNOTE 1: Código)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[17],
        new_text_before_fn="Mediante escrito del 14 de mayo de 2026, subsanado mediante escrito del 10 de junio de 2026, el señor Reyes denunció a Pacífico Seguros y a la Clínica por presuntas infracciones a la Ley N° 29571, Código de Protección y Defensa del Consumidor",
        new_text_after_fn=" (en adelante, el Código), señalando lo siguiente:"
    )
    
    # P19 a P31: Hechos
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[19],
        full_text="El 15 de marzo de 2026, el vehículo de placa de rodaje N° B8K-412 en el que se transportaba sufrió un accidente de tránsito de magnitud, contando dicha unidad con el Seguro Obligatorio de Accidentes de Tránsito – Póliza N° 0588291044 emitido por Pacífico Seguros (en adelante, SOAT)."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[21],
        full_text="A consecuencia del impacto, fue ingresado en estado de emergencia a la Clínica, donde requirió intervención quirúrgica de urgencia mediante osteosíntesis traumatológica de tibia y peroné, permaneciendo hospitalizado durante cinco (5) días calendarios."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[23],
        full_text="Pacífico Seguros se negó injustificadamente a otorgar la cobertura integral de gastos médicos del SOAT, alegando mediante comunicación denegatoria que la Clínica no formaba parte de su red preferente de atención médica para accidentes vehiculares."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[25],
        full_text="Asimismo, la Clínica emitió una preliquidación y liquidación asistencial exigiendo el cobro no pactado de S/ 6 800,00 por concepto de material médico quirúrgico de osteosíntesis no catalogado en el presupuesto inicial informado."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[27],
        full_text="Adicionalmente, al solicitar el alta y la documentación de su internamiento, la Clínica se negó a entregarle copia íntegra y fedateada de su historia clínica, condicionando su egreso hospitalario a la suscripción forzosa de un pagaré bancario en garantía por el saldo pendiente."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[29],
        full_text="Pese a los reiterados requerimientos formales efectuados por el señor Reyes, Pacífico Seguros ratificó su negativa de cobertura y la Clínica mantuvo la retención de su expediente clínico."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[31],
        full_text="Por tales motivos, interpuso denuncia administrativa ante el Indecopi solicitando medidas correctivas y sanciones para ambos proveedores en el ámbito de la relación de consumo."
    )
    
    # P33: Medidas correctivas solicitadas
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[33],
        full_text="El señor Reyes solicitó, en calidad de medida correctiva, que: (i) Pacífico Seguros cumpla con otorgar la cobertura integral de gastos médicos del SOAT; (ii) la Clínica cumpla con reintegrar el cobro indebido de S/ 6 800,00 y hacer entrega formal de la copia completa fedateada de la historia clínica; y, (iii) se declare la inexigibilidad del pagaré suscrito bajo coacción. Asimismo, solicitó la imposición de costas y costos."
    )
    
    # P37: Subtítulo ANÁLISIS
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[37],
        full_text="Sobre la improcedencia de la denuncia interpuesta por el señor Reyes"
    )
    
    # P45: Verificación de competencia
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[45],
        full_text="Al respecto, la Comisión de Protección al Consumidor N° 1 (en adelante, la Comisión) debe verificar si el Indecopi es el órgano competente para pronunciarse sobre los hechos cuestionados por el señor Reyes en su denuncia contra Pacífico Seguros y la Clínica."
    )
    
    # P59: Subsunción normativa de IAFAS + IPRESS (TIENE FOOTNOTE 9: Anexo I-C y I-B)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[59],
        new_text_before_fn="Asimismo, conforme al Anexo I-C del Reglamento de Infracciones y Sanciones de Susalud, aprobado por el Decreto Supremo N° 031-2014-SA",
        new_text_after_fn=", dicha entidad se encuentra facultada para sancionar el hecho de no brindar cobertura oportuna a los afiliados o beneficiarios respecto de las IAFAS; y, conforme al Anexo I-B del citado cuerpo normativo, sancionar a las IPRESS por realizar cobros no pactados, deficiencias asistenciales o negativa de entrega de documentación médica obligatoria."
    )
    
    # P77: Determinación de autoridad competente
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[77],
        full_text="A continuación, se determinará la autoridad competente para conocer la denuncia interpuesta por el señor Reyes contra Pacífico Seguros y la Clínica, de ser el caso, imponer las sanciones correspondientes."
    )
    
    # P85: Subsunción fáctica al caso concreto (calibrado a longitud estándar)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[85],
        full_text="En el presente caso, el señor Reyes cuestionó que: (i) Pacífico Seguros (IAFAS) no brindó oportunamente la cobertura de gastos médicos derivada del SOAT; (ii) la Clínica (IPRESS) exigió un cobro excesivo no pactado de S/ 6 800,00 por material médico; y, (iii) la Clínica se negó a entregar su historia clínica condicionando el alta. Dichas conductas inciden directamente en la cobertura del aseguramiento y en las prestaciones asistenciales de salud."
    )
    
    # P87: Conclusión de improcedencia
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[87],
        full_text="En consecuencia, la Comisión considera que corresponde declarar improcedente la denuncia interpuesta por el señor Reyes contra Pacífico Seguros y la Clínica, en la medida que las conductas cuestionadas resultan ser materia de competencia exclusiva de Susalud."
    )
    
    # P93: RESUELVE PRIMERO (TIENE FOOTNOTE 17: Devolución de tasa)
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[93],
        new_text_before_fn="PRIMERO: declarar improcedente la denuncia interpuesta por el señor Marco Antonio Reyes Salvatierra contra Pacífico Compañía de Seguros y Reaseguros S.A. y Clínica San Felipe S.A., por presunta infracción a la Ley N° 29571, Código de Protección y Defensa del Consumidor, en la medida que ha quedado acreditado que las conductas cuestionadas resultan ser materia de exclusiva competencia de la Superintendencia Nacional de Salud. En consecuencia, disponer la devolución a la parte denunciante de la tasa por derecho de trámite pagada",
        new_text_after_fn=", previa solicitud a la Unidad de Finanzas y Contabilidad del Indecopi."
    )
    
    # P97: RESUELVE TERCERO (TIENE FOOTNOTES 18, 19, 20: Recursos de apelación)
    if len(doc.paragraphs[97].runs) >= 8:
        doc.paragraphs[97].runs[0].text = "TERCERO: "
        doc.paragraphs[97].runs[1].text = "informar "
        doc.paragraphs[97].runs[2].text = "al "
        doc.paragraphs[97].runs[3].text = ""
        doc.paragraphs[97].runs[4].text = "señor "
        doc.paragraphs[97].runs[5].text = ""
        doc.paragraphs[97].runs[6].text = ""
        doc.paragraphs[97].runs[7].text = "Marco Antonio Reyes Salvatierra"
                
    # P98 y P99: Agregar Comisionados y Firma Presidenta
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[98],
        full_text="Con la intervención de los señores Comisionados: Mónica Tatiana Siverio Puycan, María de Fátima Ponce Regalado, Ernesto Alonso Calderón Burneo y Aldrin Capcha Coronado."
    )
    set_paragraph_text_preserving_footnotes(
        doc.paragraphs[99],
        full_text="MÓNICA TATIANA SIVERIO PUYCAN\nPresidenta"
    )
    
    # Guardar docx modificado
    doc.save(OUT_DOCX)
    print("DOCX modificado guardado exitosamente.")
    
    # Ahora actualizamos header1.xml para el número de expediente 0512-2026/CC1
    import tempfile
    temp_zip = Path(tempfile.mktemp(suffix=".docx"))
    with ZipFile(OUT_DOCX, "r") as zin:
        with ZipFile(temp_zip, "w") as zout:
            for item in zin.infolist():
                if item.filename == "word/header1.xml":
                    hdr = zin.read(item.filename).decode("utf-8")
                    hdr = hdr.replace("2685-2025/CC1", "0512-2026/CC1")
                    zout.writestr(item, hdr.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))
    shutil.move(temp_zip, OUT_DOCX)
    print("Header actualizado con Expediente 0512-2026/CC1.")

if __name__ == "__main__":
    generate_mixta()

"""Constructor OpenXML de alta precisión para Resoluciones de Improcedencia a SUSALUD.
Comisión de Protección al Consumidor N° 1 (CC1) - INDECOPI.

Filosofía: NUNCA generar un .docx desde la nada. Se clona la plantilla maestra institucional
preservando encabezados, pies de página, numeración y membrete oficial del Indecopi;
se limpia el cuerpo del documento y se inyecta la resolución con micro-tipografía
calibrada y formateo determinista.
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from src.config import (
    AUTORIDAD_RESOLUTIVA,
    FIRMA_RESOLUTIVA,
    FUENTE_PRINCIPAL,
    PLANTILLAS_DIR,
    GENERADOS_DIR,
    TAMANO_TEXTO_PT,
    TAMANO_TITULO_PT,
)

INDENT_ORDINAL_LEFT = Inches(0.39)
INDENT_ORDINAL_HANG = Inches(-0.39)
INDENT_INCISOS_LEFT = Inches(0.79)
INDENT_INCISOS_HANG = Inches(-0.39)


def forzar_formato_parrafo(parrafo) -> None:
    """Fuerza interlineado simple 1.0 y espaciado antes/después en 0 pt."""
    pf = parrafo.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0


def add_run(
    parrafo,
    texto: str,
    *,
    bold: bool = False,
    size=Pt(TAMANO_TEXTO_PT),
    superscript: bool = False,
) -> Any:
    """Añade un run con Arial Narrow garantizado en todos los slots OpenXML."""
    run = parrafo.add_run(texto)
    run.font.name = FUENTE_PRINCIPAL
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), FUENTE_PRINCIPAL)
    rFonts.set(qn("w:hAnsi"), FUENTE_PRINCIPAL)
    rFonts.set(qn("w:cs"), FUENTE_PRINCIPAL)
    run.font.size = size
    run.bold = bold

    if superscript:
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rPr.append(vert)

    return run


def add_texto_con_superindices(
    parrafo,
    texto: str,
    *,
    bold: bool = False,
    size=Pt(TAMANO_TEXTO_PT),
) -> None:
    """Detecta 'N°' o números ordinales como 1° y formatea la volada en superíndice."""
    partes = re.split(r"(N°|N\s*°|\b\d+°)", texto)
    for p in partes:
        if not p:
            continue
        if p.startswith("N") and "°" in p:
            add_run(parrafo, "N", bold=bold, size=size)
            add_run(parrafo, "°", bold=bold, size=size, superscript=True)
        elif re.match(r"^\d+°$", p):
            num = p[:-1]
            add_run(parrafo, num, bold=bold, size=size)
            add_run(parrafo, "°", bold=bold, size=size, superscript=True)
        else:
            add_run(parrafo, p, bold=bold, size=size)


def nuevo_parrafo(
    doc: Document,
    alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY,
    left_indent=None,
    hanging_indent=None,
):
    """Crea un párrafo con interlineado simple y sin espaciados parásitos."""
    p = doc.add_paragraph()
    p.alignment = alineacion
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    if left_indent is not None:
        pf.left_indent = left_indent
    if hanging_indent is not None:
        pf.first_line_indent = hanging_indent
    return p


def construir_resolucion_improcedencia(
    datos: Dict[str, Any],
    ruta_salida: Path,
    ruta_plantilla: Path | None = None,
) -> Path:
    """Construye una resolución de improcedencia para SUSALUD (IPRESS / IAFAS)."""
    if ruta_plantilla is None:
        plantillas = list(PLANTILLAS_DIR.glob("*.docx"))
        if not plantillas:
            raise FileNotFoundError(f"No se encontró ninguna plantilla base en {PLANTILLAS_DIR}")
        ruta_plantilla = plantillas[0]

    # Copia inicial
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ruta_plantilla, ruta_salida)

    doc = Document(str(ruta_salida))

    # Limpiar cuerpo manteniendo encabezados/pies
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    # 1. ENCABEZADO INSTITUCIONAL
    p_aut = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p_aut, AUTORIDAD_RESOLUTIVA, bold=True)

    p_num = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
    num_res = datos.get("numero_resolucion", "RESOLUCIÓN N° 0001-2026/CC1")
    add_texto_con_superindices(p_num, num_res, bold=True)

    nuevo_parrafo(doc)

    # 2. CUADRO DE IDENTIFICACIÓN
    p_exp = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_run(p_exp, "EXPEDIENTE\t: ", bold=True)
    add_run(p_exp, datos.get("expediente", ""))

    p_den = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_run(p_den, "DENUNCIANTE\t: ", bold=True)
    add_run(p_den, datos.get("denunciante", ""))

    p_denun = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_run(p_denun, "DENUNCIADO(S)\t: ", bold=True)
    add_run(p_denun, datos.get("denunciado", ""))

    p_mat = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_run(p_mat, "MATERIA\t\t: ", bold=True)
    add_run(p_mat, "INCOMPETENCIA POR RAZÓN DE LA MATERIA / SUSALUD")

    p_proc = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_run(p_proc, "PROCEDENCIA\t: ", bold=True)
    add_run(p_proc, "COMISIÓN DE PROTECCIÓN AL CONSUMIDOR N° 1")

    nuevo_parrafo(doc)

    # 3. FECHA
    p_fecha = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
    add_run(p_fecha, f"Lima, {datos.get('fecha_emision', '25 de septiembre de 2026')}.")

    nuevo_parrafo(doc)

    # 4. VISTOS
    p_vistos_tit = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
    add_run(p_vistos_tit, "VISTOS:", bold=True)

    for visto in datos.get("vistos", []):
        p_v = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_texto_con_superindices(p_v, visto)

    nuevo_parrafo(doc)

    # 5. CONSIDERANDO
    p_cons_tit = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
    add_run(p_cons_tit, "CONSIDERANDO:", bold=True)

    for seccion in datos.get("secciones_considerando", []):
        p_sub = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.LEFT)
        add_texto_con_superindices(p_sub, seccion.get("titulo", ""), bold=True)
        for parrafo_texto in seccion.get("parrafos", []):
            p_c = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
            add_texto_con_superindices(p_c, parrafo_texto)
        nuevo_parrafo(doc)

    # 6. RESUELVE
    p_resuelve_tit = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p_resuelve_tit, "RESUELVE:", bold=True)

    for articulo in datos.get("articulos_resuelve", []):
        p_art = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_texto_con_superindices(p_art, articulo)
        nuevo_parrafo(doc)

    # 7. INTERVENCIÓN DE LA COMISIÓN (ÓRGANO COLEGIADO)
    p_interv = nuevo_parrafo(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_run(
        p_interv,
        "Con la intervención de los señores comisionados: ",
        bold=False,
    )
    add_run(p_interv, datos.get("comisionados_texto", "miembros integrantes de la Comisión de Protección al Consumidor N° 1."))

    doc.save(str(ruta_salida))
    return ruta_salida
